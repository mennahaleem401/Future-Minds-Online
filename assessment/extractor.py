import csv
import io
import json
import re
from typing import List, Dict, Tuple, Any

# Map various representations of choices to standard 'A', 'B', 'C', 'D'
CHOICE_MAP = {
    'A': 'A', 'B': 'B', 'C': 'C', 'D': 'D',
    'a': 'A', 'b': 'B', 'c': 'C', 'd': 'D',
    'أ': 'A', 'ب': 'B', 'ج': 'C', 'د': 'D',
    'ا': 'A', 'ء': 'A',
    '1': 'A', '2': 'B', '3': 'C', '4': 'D',
}


def normalize_choice(val: Any) -> str:
    """Normalize choice value to 'A', 'B', 'C', or 'D'."""
    if not val:
        return 'A'
    s = str(val).strip()
    match = re.search(r'[A-Da-dأ-د1-4]', s)
    if match:
        ch = match.group(0)
        return CHOICE_MAP.get(ch, 'A')
    return 'A'


def identify_column(header: str) -> str:
    """Identify which question field a column header corresponds to."""
    if not header:
        return ''
    cleaned = re.sub(r'[_\-\:\.]', ' ', str(header).strip().lower()).strip()

    if any(k in cleaned for k in ['explanation', 'note', 'شرح', 'توضيح', 'تعليق']):
        return 'expl'
    if any(k in cleaned for k in ['correct', 'answer', 'key', 'إجابة', 'اجابة', 'الجواب', 'الحل']):
        return 'correct'
    if any(k in cleaned for k in ['question', 'prompt', 'سؤال']):
        return 'prompt'

    # Choices matching
    if cleaned in ['a', 'opt a', 'option a', 'أ', 'اختيار أ', 'الخيار أ'] or re.search(r'^(opt|option|اختيار)?\s*[aأ1]$', cleaned):
        return 'opt_a'
    if cleaned in ['b', 'opt b', 'option b', 'ب', 'اختيار ب', 'الخيار ب'] or re.search(r'^(opt|option|اختيار)?\s*[bب2]$', cleaned):
        return 'opt_b'
    if cleaned in ['c', 'opt c', 'option c', 'ج', 'اختيار ج', 'الخيار ج'] or re.search(r'^(opt|option|اختيار)?\s*[cج3]$', cleaned):
        return 'opt_c'
    if cleaned in ['d', 'opt d', 'option d', 'د', 'اختيار د', 'الخيار د'] or re.search(r'^(opt|option|اختيار)?\s*[dد4]$', cleaned):
        return 'opt_d'

    return ''


def extract_questions_from_file(uploaded_file) -> Tuple[List[Dict[str, Any]], str]:
    """
    Extract questions from an uploaded file.
    Supports: .docx, .xlsx, .csv, .txt, .json.
    Returns: (list_of_question_dicts, error_message)
    """
    filename = getattr(uploaded_file, 'name', '').lower()
    content = uploaded_file.read()

    try:
        if filename.endswith('.docx'):
            return parse_docx(content)
        elif filename.endswith('.xlsx'):
            return parse_xlsx(content)
        elif filename.endswith('.csv'):
            return parse_csv(content)
        elif filename.endswith('.json'):
            return parse_json(content)
        elif filename.endswith('.txt') or filename.endswith('.text'):
            return parse_txt(content.decode('utf-8-sig', errors='replace'))
        else:
            try:
                text = content.decode('utf-8-sig')
                return parse_txt(text)
            except Exception:
                return [], f"صيغة الملف غير مدعومة ({filename}). الصيغ المدعومة هي: Word (.docx)، Excel (.xlsx)، CSV (.csv)، Text (.txt)، و JSON."
    except Exception as e:
        return [], f"حدث خطأ أثناء قراءة الملف: {str(e)}"


def parse_txt(text: str) -> Tuple[List[Dict[str, Any]], str]:
    """
    Parse questions from formatted plain text.
    Handles various question numbering, options (A-D, أ-د), asterisks, answers, and explanations.
    """
    lines = [line.strip() for line in text.splitlines()]
    questions = []

    current_q = None
    collecting_prompt = False

    def finish_question(q):
        if not q or not q.get('prompt'):
            return
        if not q.get('option_a'):
            q['option_a'] = 'صواب (True)'
        if not q.get('option_b'):
            q['option_b'] = 'خطأ (False)'
        if not q.get('correct_option'):
            q['correct_option'] = 'A'
        q['points'] = q.get('points', 1)
        questions.append(q)

    q_num_pattern = re.compile(r'^(?:[Qq][0-9]+[:\.\-]|[0-9]+[:\.\-\)]|س\s*[0-9]+[:\.\-])\s*(.*)$')
    opt_pattern = re.compile(r'^(\*?)\s*([A-Da-dأ-د1-4])[\.\-\)\:\s]\s*(\*?)\s*(.*)$')
    ans_pattern = re.compile(r'^(?:[Aa]nswer|[Cc]orrect(?:\s+[Aa]nswer)?|[Kk]ey|الإجابة|الاجابة|الجواب|الحل)(?:\s+الصحيحة)?[\:\s\-]+(.*)$', re.IGNORECASE)
    exp_pattern = re.compile(r'^(?:[Ee]xplanation|[Nn]ote|الشرح|توضيح|ملاحظة)[\:\s\-]+(.*)$', re.IGNORECASE)

    for line in lines:
        if not line:
            continue

        ans_match = ans_pattern.match(line)
        if ans_match and current_q:
            val = ans_match.group(1).strip()
            current_q['correct_option'] = normalize_choice(val)
            collecting_prompt = False
            continue

        exp_match = exp_pattern.match(line)
        if exp_match and current_q:
            current_q['explanation'] = exp_match.group(1).strip()
            collecting_prompt = False
            continue

        opt_match = opt_pattern.match(line)
        if opt_match and current_q:
            is_star = bool(opt_match.group(1) or opt_match.group(3))
            choice_char = normalize_choice(opt_match.group(2))
            opt_text = opt_match.group(4).strip()

            if is_star:
                current_q['correct_option'] = choice_char

            if choice_char == 'A':
                current_q['option_a'] = opt_text
            elif choice_char == 'B':
                current_q['option_b'] = opt_text
            elif choice_char == 'C':
                current_q['option_c'] = opt_text
            elif choice_char == 'D':
                current_q['option_d'] = opt_text

            collecting_prompt = False
            continue

        q_match = q_num_pattern.match(line)
        if q_match:
            finish_question(current_q)
            current_q = {
                'prompt': q_match.group(1).strip() or line,
                'option_a': '',
                'option_b': '',
                'option_c': '',
                'option_d': '',
                'correct_option': 'A',
                'explanation': '',
                'points': 1,
            }
            collecting_prompt = True
            continue

        if current_q and collecting_prompt:
            current_q['prompt'] += "\n" + line
        elif not current_q:
            current_q = {
                'prompt': line,
                'option_a': '',
                'option_b': '',
                'option_c': '',
                'option_d': '',
                'correct_option': 'A',
                'explanation': '',
                'points': 1,
            }
            collecting_prompt = True

    finish_question(current_q)

    if not questions:
        return [], "لم يتم العثور على أسئلة بتنسيق واضح في الملف. يرجى التأكد من كتابة كل سؤال ومتبوعاً بالاختيارات والإجابة."
    return questions, ""


def parse_docx(content: bytes) -> Tuple[List[Dict[str, Any]], str]:
    """Parse questions from Word .docx document."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        full_text = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                full_text.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    full_text.append(" | ".join(row_cells))

        combined = "\n".join(full_text)
        return parse_txt(combined)
    except Exception as e:
        return [], f"خطأ في قراءة ملف Word: {str(e)}"


def parse_xlsx(content: bytes) -> Tuple[List[Dict[str, Any]], str]:
    """Parse questions from Excel .xlsx file."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return [], "ملف Excel فارغ."

        header_row = [str(cell).strip().lower() if cell is not None else '' for cell in rows[0]]
        start_idx = 1 if any(kw in "".join(header_row) for kw in ['question', 'سؤال', 'prompt', 'اختيار', 'option', 'إجابة']) else 0

        col_map = {'prompt': -1, 'opt_a': -1, 'opt_b': -1, 'opt_c': -1, 'opt_d': -1, 'correct': -1, 'expl': -1}
        for i, col in enumerate(header_row):
            field = identify_column(col)
            if field and col_map[field] == -1:
                col_map[field] = i

        questions = []
        for row in rows[start_idx:]:
            if not any(row):
                continue

            def get_val(idx, fallback_idx=None):
                if idx >= 0 and idx < len(row) and row[idx] is not None:
                    return str(row[idx]).strip()
                if fallback_idx is not None and fallback_idx < len(row) and row[fallback_idx] is not None:
                    return str(row[fallback_idx]).strip()
                return ''

            prompt = get_val(col_map['prompt'], 0)
            if not prompt:
                continue

            opt_a = get_val(col_map['opt_a'], 1)
            opt_b = get_val(col_map['opt_b'], 2)
            opt_c = get_val(col_map['opt_c'], 3)
            opt_d = get_val(col_map['opt_d'], 4)
            correct = get_val(col_map['correct'], 5)
            expl = get_val(col_map['expl'], 6)

            questions.append({
                'prompt': prompt,
                'option_a': opt_a or 'صواب (True)',
                'option_b': opt_b or 'خطأ (False)',
                'option_c': opt_c,
                'option_d': opt_d,
                'correct_option': normalize_choice(correct),
                'explanation': expl,
                'points': 1
            })

        if not questions:
            return [], "لم يتم العثور على أي أسئلة صالحة داخل جدول Excel."
        return questions, ""
    except Exception as e:
        return [], f"خطأ في قراءة ملف Excel: {str(e)}"


def parse_csv(content: bytes) -> Tuple[List[Dict[str, Any]], str]:
    """Parse questions from CSV file."""
    try:
        text = content.decode('utf-8-sig', errors='replace')
        reader = csv.reader(io.StringIO(text))
        rows = [r for r in reader if any(cell.strip() for cell in r)]
        if not rows:
            return [], "ملف CSV فارغ."

        header_row = [cell.strip().lower() for cell in rows[0]]
        start_idx = 1 if any(kw in "".join(header_row) for kw in ['question', 'سؤال', 'prompt', 'اختيار', 'option', 'إجابة']) else 0

        col_map = {'prompt': -1, 'opt_a': -1, 'opt_b': -1, 'opt_c': -1, 'opt_d': -1, 'correct': -1, 'expl': -1}
        for i, col in enumerate(header_row):
            field = identify_column(col)
            if field and col_map[field] == -1:
                col_map[field] = i

        questions = []
        for row in rows[start_idx:]:
            if not any(row):
                continue

            def get_val(idx, fallback_idx=None):
                if idx >= 0 and idx < len(row):
                    return row[idx].strip()
                if fallback_idx is not None and fallback_idx < len(row):
                    return row[fallback_idx].strip()
                return ''

            prompt = get_val(col_map['prompt'], 0)
            if not prompt:
                continue

            opt_a = get_val(col_map['opt_a'], 1)
            opt_b = get_val(col_map['opt_b'], 2)
            opt_c = get_val(col_map['opt_c'], 3)
            opt_d = get_val(col_map['opt_d'], 4)
            correct = get_val(col_map['correct'], 5)
            expl = get_val(col_map['expl'], 6)

            questions.append({
                'prompt': prompt,
                'option_a': opt_a or 'صواب (True)',
                'option_b': opt_b or 'خطأ (False)',
                'option_c': opt_c,
                'option_d': opt_d,
                'correct_option': normalize_choice(correct),
                'explanation': expl,
                'points': 1
            })

        if not questions:
            return [], "لم يتم العثور على أي أسئلة صالحة في ملف CSV."
        return questions, ""
    except Exception as e:
        return [], f"خطأ في قراءة ملف CSV: {str(e)}"


def parse_json(content: bytes) -> Tuple[List[Dict[str, Any]], str]:
    """Parse questions from JSON file."""
    try:
        data = json.loads(content.decode('utf-8-sig'))
        raw_list = data if isinstance(data, list) else data.get('questions', [])
        if not raw_list:
            return [], "لم يتم العثور على مصفوفة أسئلة (questions) في ملف JSON."

        questions = []
        for item in raw_list:
            prompt = item.get('prompt') or item.get('question') or item.get('سؤال')
            if not prompt:
                continue
            questions.append({
                'prompt': prompt,
                'option_a': item.get('option_a') or item.get('a') or item.get('أ') or 'صواب',
                'option_b': item.get('option_b') or item.get('b') or item.get('ب') or 'خطأ',
                'option_c': item.get('option_c') or item.get('c') or item.get('ج') or '',
                'option_d': item.get('option_d') or item.get('d') or item.get('د') or '',
                'correct_option': normalize_choice(item.get('correct_option') or item.get('answer') or item.get('إجابة')),
                'explanation': item.get('explanation') or item.get('شرح') or '',
                'points': int(item.get('points', 1))
            })

        if not questions:
            return [], "لم يتم العثور على أسئلة صالحة في ملف JSON."
        return questions, ""
    except Exception as e:
        return [], f"خطأ في قراءة ملف JSON: {str(e)}"
