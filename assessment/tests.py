from datetime import timedelta
from decimal import Decimal
from django.test import TestCase, Client
from django.urls import reverse
from django.utils import timezone
import io
from django.core.files.uploadedfile import SimpleUploadedFile
from core.models import User, StudentProfile, InstructorProfile
from academy.models import Course
from assessment.models import Assignment, Submission, Quiz, Question, QuizAttempt, StudentProgress
from assessment.progress_service import recalculate_student_progress


class AssessmentTests(TestCase):
    def setUp(self):
        self.client = Client()
        
        # Instructor & Student
        inst_user = User.objects.create_user(username='inst_user', password='password123', role=User.Role.INSTRUCTOR)
        self.instructor = InstructorProfile.objects.create(user=inst_user, specialization='Python')
        
        student_user = User.objects.create_user(username='student_user', password='password123', role=User.Role.STUDENT)
        self.student = StudentProfile.objects.create(user=student_user, learning_mode=StudentProfile.LearningMode.ONLINE)

        # Course
        self.course = Course.objects.create(title='Assessment Course', instructor=self.instructor)
        self.course.enrolled_students.add(self.student)

        # Assignment
        self.assignment = Assignment.objects.create(
            course=self.course,
            instructor=self.instructor,
            title='Test HW',
            description='Test Description',
            due_date=timezone.now() + timedelta(days=2),
            max_score=100
        )

        # Quiz
        self.quiz = Quiz.objects.create(
            course=self.course,
            instructor=self.instructor,
            title='Test Quiz',
            time_limit_minutes=10,
            passing_percentage=70
        )
        self.q1 = Question.objects.create(quiz=self.quiz, prompt='Q1', option_a='A1', option_b='B1', correct_option='A', points=1)
        self.q2 = Question.objects.create(quiz=self.quiz, prompt='Q2', option_a='A2', option_b='B2', correct_option='B', points=1)

    def test_student_can_submit_assignment_and_progress_updates(self):
        self.client.login(username='student_user', password='password123')
        submit_url = reverse('assignment_submit', kwargs={'assignment_id': self.assignment.id})
        response = self.client.post(submit_url, {'submission_text': 'print("Hello World")'})
        self.assertRedirects(response, reverse('assignment_detail', kwargs={'assignment_id': self.assignment.id}))

        sub = Submission.objects.get(assignment=self.assignment, student=self.student)
        self.assertEqual(sub.submission_text, 'print("Hello World")')
        self.assertEqual(sub.status, Submission.Status.SUBMITTED)

    def test_quiz_auto_scoring(self):
        self.client.login(username='student_user', password='password123')
        take_url = reverse('quiz_take', kwargs={'quiz_id': self.quiz.id})
        # Q1 correct (A), Q2 correct (B) -> 100%
        response = self.client.post(take_url, {
            f'q_{self.q1.id}': 'A',
            f'q_{self.q2.id}': 'B'
        })
        self.assertEqual(response.status_code, 200)

        attempt = QuizAttempt.objects.get(quiz=self.quiz, student=self.student)
        self.assertEqual(attempt.score, 2)
        self.assertEqual(attempt.percentage, Decimal('100.00'))
        self.assertTrue(attempt.passed)

        # Recalculate progress
        progress = recalculate_student_progress(self.student, self.course)
        self.assertGreater(progress.overall_percentage, 0)

    def test_student_blocked_from_retaking_quiz(self):
        self.client.login(username='student_user', password='password123')
        take_url = reverse('quiz_take', kwargs={'quiz_id': self.quiz.id})

        # First attempt (allowed)
        resp1 = self.client.post(take_url, {
            f'q_{self.q1.id}': 'A',
            f'q_{self.q2.id}': 'B'
        })
        self.assertEqual(resp1.status_code, 200)
        self.assertEqual(QuizAttempt.objects.filter(quiz=self.quiz, student=self.student).count(), 1)

        # Second attempt via GET (blocked)
        resp_get = self.client.get(take_url)
        self.assertEqual(resp_get.status_code, 200)
        self.assertTrue(resp_get.context.get('already_submitted'))
        self.assertIn('لقد قمت بحل هذا الكويز مسبقاً', resp_get.content.decode('utf-8'))

        # Second attempt via POST (blocked)
        resp_post = self.client.post(take_url, {
            f'q_{self.q1.id}': 'B',
            f'q_{self.q2.id}': 'A'
        })
        self.assertEqual(resp_post.status_code, 200)
        self.assertTrue(resp_post.context.get('already_submitted'))
        self.assertEqual(QuizAttempt.objects.filter(quiz=self.quiz, student=self.student).count(), 1)

    def test_extract_questions_from_txt(self):
        from assessment.extractor import extract_questions_from_file
        txt_content = (
            "1. ما هو المتغير في بايثون؟\n"
            "A) مكان لتخزين البيانات\n"
            "B) دالة تشغيل\n"
            "C) ملف نظام\n"
            "D) لا شيء مما سبق\n"
            "الإجابة: A\n"
            "الشرح: المتغيرات تستخدم لتخزين القيم في الذاكرة.\n\n"
            "Q2. What is 5 + 5?\n"
            "A) 8\n"
            "B) 10\n"
            "Answer: B\n"
        ).encode('utf-8')

        uploaded = SimpleUploadedFile("quiz_questions.txt", txt_content, content_type="text/plain")
        questions, err = extract_questions_from_file(uploaded)
        self.assertEqual(err, "")
        self.assertEqual(len(questions), 2)
        self.assertEqual(questions[0]['correct_option'], 'A')
        self.assertEqual(questions[1]['correct_option'], 'B')

    def test_extract_questions_from_csv(self):
        from assessment.extractor import extract_questions_from_file
        csv_content = (
            "السؤال,الاختيار أ,الاختيار ب,الاختيار ج,الاختيار د,الإجابة الصحيحة,الشرح\n"
            "ما هي دالة الطباعة؟,print,echo,write,display,A,دالة الطباعة في بايثون\n"
            "هل بايثون لغة مفسرة؟,نعم,لا,,,أ,بايثون لغة interpreted\n"
        ).encode('utf-8-sig')

        uploaded = SimpleUploadedFile("quiz.csv", csv_content, content_type="text/csv")
        questions, err = extract_questions_from_file(uploaded)
        self.assertEqual(err, "")
        self.assertEqual(len(questions), 2)
        self.assertEqual(questions[0]['prompt'], "ما هي دالة الطباعة؟")
        self.assertEqual(questions[0]['correct_option'], "A")
        self.assertEqual(questions[1]['correct_option'], "A")

    def test_extract_questions_from_docx(self):
        from assessment.extractor import extract_questions_from_file
        from docx import Document
        doc = Document()
        doc.add_paragraph("1. ما هي الكلمة المحجوزة لتعريف دالة؟")
        doc.add_paragraph("A) function")
        doc.add_paragraph("B) def")
        doc.add_paragraph("C) fun")
        doc.add_paragraph("D) define")
        doc.add_paragraph("الإجابة: B")
        doc.add_paragraph("الشرح: كلمة def تستخدم لتعريف الدوال.")

        bio = io.BytesIO()
        doc.save(bio)
        bio.seek(0)

        uploaded = SimpleUploadedFile("exam.docx", bio.read(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        questions, err = extract_questions_from_file(uploaded)
        self.assertEqual(err, "")
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0]['correct_option'], "B")
        self.assertEqual(questions[0]['option_b'], "def")

    def test_extract_questions_from_xlsx(self):
        from assessment.extractor import extract_questions_from_file
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Question", "Option A", "Option B", "Option C", "Option D", "Correct", "Explanation"])
        ws.append(["What is type of 5?", "int", "float", "str", "bool", "A", "5 is integer"])

        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)

        uploaded = SimpleUploadedFile("questions.xlsx", bio.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        questions, err = extract_questions_from_file(uploaded)
        self.assertEqual(err, "")
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0]['correct_option'], "A")
        self.assertEqual(questions[0]['option_a'], "int")

    def test_upload_questions_view(self):
        self.client.login(username='inst_user', password='password123')
        upload_url = reverse('quiz_upload_questions', kwargs={'quiz_id': self.quiz.id})

        txt_content = (
            "1. سؤال اختباري جديد؟\n"
            "A) اختيار 1\n"
            "B) اختيار 2\n"
            "Answer: A\n"
        ).encode('utf-8')
        uploaded = SimpleUploadedFile("new_q.txt", txt_content, content_type="text/plain")

        resp = self.client.post(upload_url, {
            'questions_file': uploaded,
            'replace_existing': '0'
        })
        self.assertRedirects(resp, reverse('quiz_detail', kwargs={'quiz_id': self.quiz.id}))
        # Initially had 2 questions, now should have 3
        self.assertEqual(self.quiz.questions.count(), 3)

    def test_question_edit_view(self):
        self.client.login(username='inst_user', password='password123')
        edit_url = reverse('question_edit', kwargs={'question_id': self.q1.id})

        resp = self.client.post(edit_url, {
            'prompt': 'Updated Question 1 Prompt',
            'option_a': 'New Opt A',
            'option_b': 'New Opt B',
            'option_c': 'New Opt C',
            'option_d': 'New Opt D',
            'correct_option': 'C',
            'explanation': 'New Explanation',
            'points': 2
        })
        self.assertRedirects(resp, reverse('quiz_detail', kwargs={'quiz_id': self.quiz.id}))

        self.q1.refresh_from_db()
        self.assertEqual(self.q1.prompt, 'Updated Question 1 Prompt')
        self.assertEqual(self.q1.option_a, 'New Opt A')
        self.assertEqual(self.q1.correct_option, 'C')
        self.assertEqual(self.q1.explanation, 'New Explanation')
        self.assertEqual(self.q1.points, 2)

    def test_question_delete_view(self):
        self.client.login(username='inst_user', password='password123')
        delete_url = reverse('question_delete', kwargs={'question_id': self.q2.id})

        resp = self.client.post(delete_url)
        self.assertRedirects(resp, reverse('quiz_detail', kwargs={'quiz_id': self.quiz.id}))
        self.assertFalse(Question.objects.filter(id=self.q2.id).exists())
        self.assertEqual(self.quiz.questions.count(), 1)

    def test_question_add_single_view(self):
        self.client.login(username='inst_user', password='password123')
        add_url = reverse('question_add_single', kwargs={'quiz_id': self.quiz.id})

        resp = self.client.post(add_url, {
            'prompt': 'Manually Added Question',
            'option_a': 'Alpha',
            'option_b': 'Beta',
            'correct_option': 'B',
            'points': 3
        })
        self.assertRedirects(resp, reverse('quiz_detail', kwargs={'quiz_id': self.quiz.id}))
        new_q = self.quiz.questions.filter(prompt='Manually Added Question').first()
        self.assertIsNotNone(new_q)
        self.assertEqual(new_q.correct_option, 'B')
        self.assertEqual(new_q.points, 3)

